"""WAFCT 2019 food composition importer."""

from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from aya_api.models import Food

SOURCE = "WAFCT"
SOURCE_VERSION = "2019"


async def import_wafct(session: AsyncSession, xlsx_path: Path) -> int:
    """Import WAFCT 2019 workbook into food table."""
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["03 NV_sum_39 (per 100g EP)"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip() if c is not None else "" for c in rows[1]]
    count = 0
    for row in rows[2:]:
        if not row or row[0] is None:
            continue
        code = str(row[0]).strip()
        name = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        nutrients = {}
        for i, h in enumerate(header):
            if i < 2 or not h:  # noqa: PLR2004
                continue
            v = row[i] if i < len(row) else None
            if v is not None:
                nutrients[h] = v
        session.add(
            Food(
                food_code=code,
                name=name,
                nutrients=nutrients,
                source=SOURCE,
                source_version=SOURCE_VERSION,
            )
        )
        count += 1
    await session.commit()
    return count
